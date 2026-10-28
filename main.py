import os
import csv
from colorama import init, Fore
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from procesamiento import limpiar_y_filtrar_lecturas, calcular_metricas_produccion

init(autoreset=True)

carpeta = "Datos_Planta"
os.makedirs(carpeta, exist_ok=True)  # crea la carpeta si no existe

log_crudo = ["105", "98.5", "abc", "nan", "120", "inf", "87", "", "110.2", "ERROR_SENSOR"]

validos, incidencias = limpiar_y_filtrar_lecturas(log_crudo)

ruta_bitacora = os.path.join(carpeta, "bitacora_alertas.txt")
with open(ruta_bitacora, "a", encoding="utf-8") as archivo:
    for incidencia in incidencias:
        archivo.write(incidencia + "\n")

print(Fore.YELLOW + f"Se registraron {len(incidencias)} incidencias en la bitácora.")

minimo, maximo, promedio = calcular_metricas_produccion(validos)

ruta_csv = os.path.join(carpeta, "reporte_metricas.csv")
with open(ruta_csv, "w", newline="", encoding="utf-8-sig") as archivo:
    escritor = csv.writer(archivo, delimiter=";")
    escritor.writerow(["Mínimo", "Máximo", "Promedio"])
    escritor.writerow([minimo, maximo, promedio])

print(Fore.GREEN + "Reporte CSV generado correctamente.")

wb = Workbook()
ws = wb.active
ws.title = "Producción_Mensual"

ws["A1"] = "Código Componente"
ws["B1"] = "Cantidad Producida"

fuente = Font(bold=True, color="FFFFFF")
relleno = PatternFill(start_color="000080", end_color="000080", fill_type="solid")
for celda in (ws["A1"], ws["B1"]):
    celda.font = fuente
    celda.fill = relleno

datos = [
    ("COMP-001", 500),
    ("COMP-002", 750),
    ("COMP-003", 620),
    ("COMP-004", 810),
    ("COMP-005", 430),
]
for fila in datos:
    ws.append(fila)

ruta_excel = os.path.join(carpeta, "Reporte_Calidad.xlsx")
wb.save(ruta_excel)
print(Fore.GREEN + "Excel creado.")

wb2 = load_workbook(ruta_excel)
ws2 = wb2.active
print(Fore.CYAN + f"Valor original de B2: {ws2['B2'].value}")
ws2["B2"] = 520  # valor corregido por control de calidad
wb2.save(ruta_excel)
print(Fore.CYAN + f"Valor corregido de B2: {ws2['B2'].value}")